"""Create a Word report from a completed CrewAI run without calling any agents."""

import json
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt


def _add_inline(paragraph, token):
    """Preserve emphasis and show link URLs in the exported document."""
    bold = False
    italic = False
    link_url = None

    for child in token.children or []:
        if child.type == "strong_open":
            bold = True
        elif child.type == "strong_close":
            bold = False
        elif child.type == "em_open":
            italic = True
        elif child.type == "em_close":
            italic = False
        elif child.type == "link_open":
            link_url = child.attrGet("href")
        elif child.type == "link_close":
            if link_url:
                paragraph.add_run(f" ({link_url})")
            link_url = None
        elif child.type in {"softbreak", "hardbreak"}:
            paragraph.add_run("\n")
        elif child.type in {"text", "code_inline"}:
            run = paragraph.add_run(child.content)
            run.bold = bold
            run.italic = italic
            if child.type == "code_inline":
                run.font.name = "Consolas"


def _add_markdown(document, text):
    """Render Markdown headings, paragraphs, lists, and tables."""
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(text or "")
    paragraph = None
    lists = []
    table = None
    row = None
    column = 0

    for token in tokens:
        kind = token.type

        if kind == "heading_open":
            level = min(int(token.tag[1:]) + 1, 4)
            paragraph = document.add_heading(level=level)
        elif kind == "bullet_list_open":
            lists.append("List Bullet")
        elif kind == "ordered_list_open":
            lists.append("List Number")
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "paragraph_open":
            paragraph = document.add_paragraph(
                style=lists[-1] if lists else None
            )
        elif kind == "table_open":
            table = document.add_table(rows=0, cols=0)
            table.style = "Table Grid"
        elif kind == "tr_open":
            row = table.add_row()
            column = 0
        elif kind in {"th_open", "td_open"}:
            if column >= len(table.columns):
                table.add_column(Inches(1.5))
            paragraph = row.cells[column].paragraphs[0]
            column += 1
        elif kind == "inline" and paragraph is not None:
            _add_inline(paragraph, token)
        elif kind == "table_close":
            table = None
            paragraph = None
        elif kind in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(token.content)
            run.font.name = "Consolas"
            run.font.size = Pt(8)


def build_word_report(
    *,
    brief,
    strategy,
    analysis,
    research_notes,
    source_checks,
    run_id,
    elapsed_seconds,
):
    """Return DOCX bytes containing the result and its research appendix."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    document.add_heading("Market Research & GTM Plan", 0)
    document.add_paragraph(
        "Provisional research output. Evidence gaps and proposed "
        "validation targets should be reviewed before business use."
    )

    document.add_heading("Research brief", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"

    fields = [
        ("Topic", brief["topic"]),
        ("Geography", brief["geography"]),
        ("Target customer", brief["target_customer"]),
        ("Implementation", "CrewAI"),
        ("Run ID", run_id),
        ("Elapsed time", f"{elapsed_seconds:.1f} seconds"),
    ]
    for label, value in fields:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

    document.add_heading("Provisional GTM strategy", level=1)
    _add_markdown(document, strategy)

    document.add_heading("Market Analyst assessment", level=1)
    _add_markdown(document, analysis)

    document.add_page_break()
    document.add_heading("Research candidates", level=1)
    try:
        findings = json.loads(research_notes)
    except (TypeError, ValueError):
        document.add_paragraph(research_notes or "No research output.")
    else:
        for index, finding in enumerate(findings, start=1):
            document.add_heading(
                f"Q{index}: {finding.get('question', '')}", level=2
            )
            document.add_paragraph(
                f"Search query: {finding.get('query', 'Not recorded')}"
            )
            document.add_paragraph(
                f"Search status: {finding.get('search_status', 'unknown')}"
            )
            document.add_paragraph(
                f"Evidence gap: {finding.get('gap_reason', '')}"
            )
            for source in finding.get("sources", []):
                document.add_paragraph(
                    f"{source.get('title', '')}\n"
                    f"{source.get('url', '')}\n"
                    f"Relevance: {source.get('relevance', 'unknown')}\n"
                    f"{source.get('snippet_or_note', '')}"
                )

    document.add_heading("Source checks and retrieved passages", level=1)
    for check in source_checks:
        document.add_heading(
            check.get("candidate_title") or "Source", level=2
        )
        document.add_paragraph(check.get("url", ""))
        document.add_paragraph(
            f"HTTP status: {check.get('http_status', 'unknown')} | "
            f"Accessible: {check.get('accessible', False)}"
        )
        document.add_paragraph(
            check.get("text_excerpt")
            or check.get("error")
            or "No readable passage extracted."
        )

    output = BytesIO()
    document.save(output)
    return output.getvalue()